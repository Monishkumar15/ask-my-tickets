"""
Step 3b: Extract plain text from documents of different file types, so
chunk.py can treat them all the same way regardless of format.
"""

import os


def load_txt(filepath):
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()


def load_md(filepath):
    # Markdown is treated as plain text -- headings/lists still read fine
    # through the existing paragraph/sentence splitter.
    with open(filepath, "r", encoding="utf-8") as f:
        return f.read()


def load_docx(filepath):
    import docx  # python-docx

    doc = docx.Document(filepath)
    blocks = [p.text.strip() for p in doc.paragraphs if p.text.strip()]

    # Also pull in table content -- ticket data can live in tables, not just
    # paragraphs. Each row becomes one pipe-joined line.
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                blocks.append(" | ".join(cells))

    return "\n\n".join(blocks)


def load_pdf(filepath):
    from pypdf import PdfReader

    reader = PdfReader(filepath)
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(p for p in pages if p.strip())
    # Scanned/image-only PDFs yield "" per page -> empty overall text ->
    # the caller (chunk.py) skips the file with a warning. No OCR support
    # by design -- out of scope for this project.


def load_xlsx(filepath):
    import openpyxl

    wb = openpyxl.load_workbook(filepath, data_only=True)
    lines = []

    for sheet in wb.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        # First row is treated as the header -- each other row becomes a
        # "Column: value" sentence-like line, which reads naturally through
        # the sentence-based chunker (unlike raw pipe-joined cell values).
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        lines.append(f"Sheet: {sheet.title}")
        for row in rows[1:]:
            pairs = [
                f"{col}: {val}"
                for col, val in zip(header, row)
                if col and val is not None and str(val).strip()
            ]
            if pairs:
                lines.append(". ".join(pairs) + ".")

    return "\n\n".join(lines)


# Dispatch by file extension -- add one line here + one function above to
# support a new format later.
LOADERS = {
    ".txt": load_txt,
    ".md": load_md,
    ".docx": load_docx,
    ".pdf": load_pdf,
    ".xlsx": load_xlsx,
}


def load_document_text(filepath):
    """
    Return extracted text for one file, or None if unreadable. Never raises
    -- the caller decides how to warn/skip.
    """
    ext = os.path.splitext(filepath)[1].lower()
    loader = LOADERS.get(ext)
    if loader is None:
        return None
    try:
        return loader(filepath)
    except Exception as e:
        print(f"WARNING: could not read '{filepath}' ({e}). Skipping.")
        return None
